"""Generate HiddenStay long-form project proposal DOCX — explanatory prose, simple English."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(r"D:\Development\Projects\sample\marketing\HiddenStay_Project_Proposal.docx")


def style_run(run, size=12, bold=False, italic=False, color=None):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color


def h(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x15, 0x10)
        run.font.name = "Calibri"


def p(doc, text, bold=False, italic=False, size=12, center=False):
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    style_run(run, size=size, bold=bold, italic=italic)
    para.paragraph_format.space_after = Pt(10)
    para.paragraph_format.line_spacing = 1.15
    return para


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(item, style="List Bullet")
        for run in para.runs:
            style_run(run, size=12)
        para.paragraph_format.space_after = Pt(4)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = header
        for para in cell.paragraphs:
            for run in para.runs:
                style_run(run, size=11, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            cell.text = str(val)
            for para in cell.paragraphs:
                for run in para.runs:
                    style_run(run, size=11)
    doc.add_paragraph()


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.9)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)

    # Cover
    p(doc, "PROJECT PROPOSAL", bold=True, size=14, center=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("HiddenStay")
    style_run(r, size=28, bold=True, color=RGBColor(0xC4, 0x5C, 0x3E))

    p(
        doc,
        "Fair-Commission Homestay Marketplace\nwith AI Trip Planner for Southeast Asia",
        bold=True,
        size=14,
        center=True,
    )
    p(
        doc,
        "A full written proposal for the capstone project\n"
        "Course: BU3102 / CP3102\n"
        "Team: Group 63 I\n"
        "James Cook University Singapore\n"
        "July 2026",
        size=12,
        center=True,
    )
    p(
        doc,
        "Live demo: https://hidden-stay.vercel.app\n"
        "This document explains the whole project in plain language so every teammate, "
        "lecturer, and reader can understand what we built and why.",
        italic=True,
        size=11,
        center=True,
    )

    doc.add_page_break()

    # 1
    h(doc, "1. Executive Summary", 1)
    p(
        doc,
        "HiddenStay is a web application for travel in Southeast Asia. It is designed for "
        "two groups of people at the same time: travellers who want authentic places to stay, "
        "and independent hosts who own small homestays, guesthouses, or boutique lodges.",
    )
    p(
        doc,
        "Today, many small hosts depend on big booking websites such as Booking.com or Expedia. "
        "Those websites are useful because they bring guests, but they often take a large share "
        "of each booking — commonly around 15% to 30%. For a small business, that fee is painful. "
        "At the same time, travellers usually plan a trip in one place (Google, maps, ChatGPT) "
        "and then book somewhere else. Planning and booking feel disconnected.",
    )
    p(
        doc,
        "HiddenStay brings these pieces together. Hosts can list for free and pay from 5% only "
        "when a guest actually books, so they keep 95% of the stay value on the base fee. "
        "Travellers can search stays, talk to an AI trip planner, see the plan on a map, and "
        "book a real HiddenStay listing in the same product. We are not trying to replace every "
        "hotel chain in the world. We focus on “hidden” independent stays that big platforms "
        "often treat as just another listing.",
    )
    p(
        doc,
        "For our capstone, we built and deployed a working MVP (minimum viable product). "
        "Anyone can open the live demo, choose a role (traveller, host, or admin), and try the "
        "main flows. This proposal tells the full story: the problem, the goals, the solution, "
        "the money model, the technology, the team, the risks, and what we have achieved so far.",
    )

    # 2
    h(doc, "2. Project Title and Basic Information", 1)
    p(
        doc,
        "The official name of the product is HiddenStay. It is a capstone MVP delivered as a "
        "website (not a native phone app). The work belongs to Group 63 I under the BU3102 / "
        "CP3102 courses at James Cook University Singapore. The live website used for demos is "
        "https://hidden-stay.vercel.app. Our focus market is Southeast Asia, especially "
        "independent homestays rather than large international hotel brands.",
    )
    table(
        doc,
        ["Item", "Detail"],
        [
            ["Project name", "HiddenStay"],
            ["Document type", "Overall project proposal"],
            ["Course", "BU3102 / CP3102 Capstone"],
            ["Team", "Group 63 I"],
            ["Institution", "James Cook University Singapore"],
            ["Live website", "https://hidden-stay.vercel.app"],
            ["Date", "July 2026"],
        ],
    )

    # 3
    h(doc, "3. Background and Problem Statement", 1)
    p(
        doc,
        "To understand why HiddenStay exists, it helps to look at both sides of the market: "
        "the host side and the guest side.",
    )

    h(doc, "3.1 The host problem", 2)
    p(
        doc,
        "Independent hosts often need online bookings to fill empty rooms. Online Travel "
        "Agencies (OTAs) are the usual channel. An OTA is a booking platform that lists many "
        "properties and charges a commission on each reservation. For independents, that "
        "commission is frequently high. If a guest pays SGD 100 and the OTA takes 18%, the "
        "host keeps only about SGD 82. On HiddenStay’s base model, the same host would keep "
        "SGD 95 after a 5% fee. Over many bookings, that difference matters a lot for a "
        "family-run lodge or village homestay.",
    )
    p(
        doc,
        "Hosts also feel that guest relationships and pricing power sit with the big platforms. "
        "They want a fairer option that is still simple to use — without paying a monthly fee "
        "just to list rooms and get paid.",
    )

    h(doc, "3.2 The traveller problem", 2)
    p(
        doc,
        "Travellers who want authentic local stays often jump between tools. They may ask an "
        "AI chatbot for ideas, open Google Maps for locations, read blogs for tips, and finally "
        "open Booking.com or Airbnb to reserve a room. That process takes time. Generic AI "
        "suggestions can also be weak: they list famous places, but they do not always connect "
        "to a stay you can book on a map next to your itinerary.",
    )

    h(doc, "3.3 Problem statement (one clear sentence)", 2)
    p(
        doc,
        "Independent hotels and homestays in Southeast Asia lose too much margin to high OTA "
        "commissions, while travellers lack one product that connects AI-assisted, map-ready "
        "trip planning to bookable local stays.",
        italic=True,
    )

    # 4
    h(doc, "4. Project Objectives", 1)
    p(
        doc,
        "We set clear objectives for the capstone so the whole team knows what “done” means. "
        "First, we wanted a working marketplace where a guest can search by city, open a "
        "property page, and complete a booking flow. Second, we wanted an AI trip planner that "
        "feels useful: the guest can chat about a trip, see days and stops on a map, and find "
        "stays that belong to our inventory. Third, we wanted a free host portal so partners "
        "can list property details, follow bookings, and understand earnings after the platform "
        "fee. Fourth, we wanted honest pricing communication — from 5% per booking, free "
        "Starter portal, and optional upgrades shown clearly, not hidden. Fifth, we wanted a "
        "live deployment so lecturers and classmates can try the product without installing "
        "software on their own computers.",
    )
    p(doc, "In short, the objectives are:", bold=True)
    bullets(
        doc,
        [
            "Prove that a fair-commission marketplace can be built and demonstrated end to end.",
            "Prove that AI planning can sit inside the same product as booking, not as a separate toy chatbot.",
            "Give hosts a usable free workspace for listings and earnings.",
            "Keep fees transparent so “5%” does not feel like a marketing trick.",
            "Ship a live demo suitable for presentation, UAT, and poster day.",
        ],
    )

    # 5
    h(doc, "5. Proposed Solution — What HiddenStay Is", 1)
    p(
        doc,
        "HiddenStay is one product with three connected experiences. Think of it as a bridge "
        "between planning and booking, with fair economics for hosts.",
    )

    h(doc, "5.1 Marketplace for guests", 2)
    p(
        doc,
        "Guests can browse and search independent stays, open a property detail page, and go "
        "through checkout. Payments in the capstone use Stripe in test (sandbox) mode, which "
        "means we can show a real payment flow without charging live cards in production. "
        "After booking, guests can see their trips. The marketplace is the part that creates "
        "the actual business event: a reservation.",
    )

    h(doc, "5.2 AI trip planner", 2)
    p(
        doc,
        "The AI planner is not meant to be a random ChatGPT clone. It is a conversion surface. "
        "A traveller describes what they want — for example a few days in Chiang Mai on a "
        "budget — and the system replies with a conversational itinerary. The plan can be shown "
        "on a map. Importantly, suggestions can connect to approved HiddenStay properties, so "
        "the guest can move from “I like this plan” to “I can book this stay nearby.” That "
        "loop is our product bet: planning intent should turn into bookings on our platform.",
    )

    h(doc, "5.3 Host portal", 2)
    p(
        doc,
        "Hosts get a free Starter portal. They can manage listings, view bookings, and see "
        "earnings after the 5% platform fee. We treat this as essential, not a paid upgrade. "
        "If a host had to pay monthly just to see bookings or get paid, our “5%” story would "
        "feel dishonest. Optional Growth tools, Featured placement, and performance ads exist "
        "for hosts who want more visibility later — but they are optional and should always be "
        "disclosed up front.",
    )

    h(doc, "5.4 Admin trust layer", 2)
    p(
        doc,
        "An admin role can approve listings before they appear to travellers. This is a basic "
        "trust step for a marketplace: guests should not book unverified junk listings during "
        "the pilot.",
    )

    h(doc, "5.5 How users enter the demo", 2)
    p(
        doc,
        "For the capstone demo, the login page lets a person choose Traveller, Host, or Admin. "
        "This is demo authentication for class use. It is not the same as a finished production "
        "system with real passwords and hardened API security. We accept that limitation for "
        "the course timeline, and we call it out honestly in this proposal.",
    )

    # 6
    h(doc, "6. Scope of Work", 1)
    h(doc, "6.1 What is included in this capstone", 2)
    p(
        doc,
        "Within the project timeline we delivered the core guest journey (search, detail, "
        "checkout, trips), the AI planner with map support, the host workspace for listings and "
        "earnings views, admin listing approval, and a public deployment so the product can be "
        "opened from a browser or QR code. We also prepared marketing and financial planning "
        "documents so the team can explain the business model, not only the code.",
    )

    h(doc, "6.2 What is not included yet", 2)
    p(
        doc,
        "We did not build native iOS or Android apps. We did not finish bank-grade production "
        "authentication for every API. We are not claiming we are ready to serve thousands of "
        "concurrent users on free hosting. We also do not try to launch every SEA country at "
        "once; the smart path is one city first (for example Chiang Mai or Da Nang), learn, "
        "then expand. Live full-scale payment operations and full payout automation are "
        "post-validation work.",
    )

    # 7
    h(doc, "7. Business Model — How HiddenStay Earns Money", 1)
    p(
        doc,
        "A common mistake in student marketplace pitches is to say “we only take 5%” and stop "
        "there. At real Southeast Asia homestay prices — often roughly SGD 35 to 55 per night, "
        "with multi-night stays around SGD 100 to 120 total — a 5% fee is only a few dollars "
        "per booking. That alone cannot pay servers, maps, AI, and marketing. So our model is "
        "designed with four streams, and we explain them together.",
    )
    p(
        doc,
        "Stream one is the booking fee: 5% of the guest total on completed bookings. This is "
        "required for all hosts. Stream two is Growth tools at about SGD 50 per month for "
        "optional SEO and analytics-style upgrades. Stream three is Featured listing at about "
        "SGD 99 per month for stronger placement when we have guest traffic. Stream four is "
        "performance advertising: an extra percentage only on bookings that come from our "
        "promotion. In slow months, ad fees can be zero. Across streams, we keep a hard idea "
        "of fairness: total fees should not exceed 15% on a booking, and Starter listing tools "
        "stay free.",
    )
    table(
        doc,
        ["Stream", "Price", "Required?"],
        [
            ["Booking fee", "5% of guest booking total", "Yes"],
            ["Growth tools", "SGD 50 / month", "Optional"],
            ["Featured listing", "SGD 99 / month per property", "Optional"],
            [
                "Performance ads",
                "7–10% on bookings we attribute to our promotion",
                "Optional · capped so total stays within 15%",
            ],
        ],
    )
    p(
        doc,
        "Worked example in plain words: a guest books a stay for SGD 100. HiddenStay takes "
        "SGD 5. The host keeps SGD 95. If the same stay had been booked through an OTA at 18%, "
        "the host might keep only about SGD 82. That story is easy for hosts to understand, "
        "and it is the core of our supply-side pitch. We still remind readers that optional "
        "upgrades exist, so nobody is surprised later.",
    )

    # 8
    h(doc, "8. Target Users and Market Focus", 1)
    p(
        doc,
        "Our primary demand-side user is an experience-led traveller who wants authentic places "
        "in Southeast Asia and does not want to juggle five apps to plan a short trip. Our "
        "primary supply-side user is an independent host — for example an eco-lodge or "
        "guesthouse owner — who currently pays high OTA fees and wants a simpler, fairer "
        "channel. Our internal user is the platform admin who protects trust by approving "
        "listings.",
    )
    p(
        doc,
        "We do not claim the entire travel market on day one. The realistic approach is one "
        "city first, build a small set of good listings, prove bookings, then consider a second "
        "city. Segment pricing for modelling sits in the homestay / guesthouse range, not luxury "
        "resort prices. That keeps our financial story honest.",
    )

    # 9
    h(doc, "9. System Design and Technology (Explained Simply)", 1)
    p(
        doc,
        "Even non-technical readers should understand how the system is put together. The "
        "website the user sees is the frontend. It is built with React (Create React App) and "
        "hosted on Vercel. When the website needs data — login, listings, chat, bookings — it "
        "calls our backend API. The backend is written with FastAPI in Python and hosted on "
        "Render. The database that stores properties, users, and bookings is MongoDB, hosted on "
        "MongoDB Atlas in the cloud. Payments go through Stripe in test mode. The AI planner "
        "uses a cloud language-model provider (Groq by default). Maps use Google Maps.",
    )
    p(
        doc,
        "A practical demo note: on Render’s free plan, the API can “sleep” after roughly fifteen "
        "minutes with no traffic. The next visitor then waits while the server wakes up. That "
        "first request can feel slow (sometimes tens of seconds). After the server is awake, "
        "normal requests are usually much faster. This is a hosting limitation of the free tier, "
        "not proof that the product idea is broken. If HiddenStay continues as a real business, "
        "paid always-on hosting is one of the first upgrades.",
    )
    table(
        doc,
        ["Layer", "Technology", "Hosting"],
        [
            ["Website (frontend)", "React", "Vercel"],
            ["API (backend)", "FastAPI (Python)", "Render"],
            ["Database", "MongoDB", "MongoDB Atlas"],
            ["Payments", "Stripe sandbox", "Stripe"],
            ["AI", "Groq (and similar)", "Cloud API"],
            ["Maps", "Google Maps", "Google Cloud"],
        ],
    )

    # 10
    h(doc, "10. Current Project Status", 1)
    p(
        doc,
        "As of this proposal, HiddenStay is past the “idea only” stage. The team has a live "
        "demo where guests can search and book in test mode, use the AI planner with map "
        "context, and where hosts and admins can use their portals. Marketing materials such as "
        "the poster content brief and financial planning docs support presentation and judging. "
        "At the same time, we are clear about unfinished production work: real password "
        "security, heavy scale, and always-on infrastructure are next-step items if the product "
        "continues after the course.",
    )
    table(
        doc,
        ["Area", "Status today"],
        [
            ["Live website demo", "Available"],
            ["Guest search and booking (test payment)", "Working"],
            ["AI trip planner and map", "Working"],
            ["Host portal", "Working"],
            ["Admin listing approval", "Working"],
            ["Poster / pitch content", "In progress / prepared for designer"],
            ["Thousands of concurrent users on free hosting", "Not ready yet"],
            ["Full production-grade authentication", "Not finished (demo login for class)"],
        ],
    )

    # 11
    h(doc, "11. Team Organisation", 1)
    p(
        doc,
        "HiddenStay is delivered by Group 63 I. Roles are split so product, engineering, "
        "design validation, analysis, and go-to-market each have an owner. In practice, "
        "everyone collaborates, but clear ownership reduces confusion during a short capstone.",
    )
    table(
        doc,
        ["Role", "Name", "Focus"],
        [
            ["Project Manager", "Zaw Latt Naung", "Coordination, scope, delivery"],
            ["Developer", "Oakkar Phyoe", "Marketplace, booking, technical build"],
            ["Design / UAT", "Eka Dian Tara Tiu", "User experience checks and testing"],
            ["BA / AI prompts", "Chenyang Gu", "Requirements and AI planner quality"],
            ["Marketing / pilots", "Yihua Deng", "Host outreach and messaging"],
        ],
    )

    # 12
    h(doc, "12. Expected Outcomes and Success Measures", 1)
    p(
        doc,
        "For the university capstone, success means we can demonstrate a believable end-to-end "
        "product: a guest can plan and book; a host can list and see earnings logic; an admin "
        "can approve supply; and outsiders can open a live URL. For the business idea, success "
        "means hosts understand the fee story, travellers understand the planner-to-book value, "
        "and the team has a realistic plan to continue with one city and careful costs rather "
        "than fantasy growth slides.",
    )
    p(
        doc,
        "Related money documents describe planning scenarios (conservative, base, optimistic) "
        "with break-even thinking over roughly an eighteen-month horizon and lean running costs. "
        "Those numbers are estimates for planning and teaching, not audited company accounts. "
        "We keep that disclaimer so readers are not misled.",
    )

    # 13
    h(doc, "13. Risks, Challenges, and Mitigations", 1)
    p(
        doc,
        "Every marketplace faces chicken-and-egg risk: without hosts, guests leave; without "
        "guests, hosts leave. Our mitigation is to start narrow (one city), keep host onboarding "
        "cheap (free portal), and use the AI planner as a reason for travellers to enter the "
        "product. Another risk is thin unit economics at 5%. We mitigate by being honest about "
        "optional upgrades and by designing ad fees that can be zero when volume is low, with a "
        "total fee cap for fairness. Technical risk includes free-tier cold starts and demo "
        "authentication; we mitigate by explaining them in demos and planning paid hosting and "
        "stronger auth if the project continues. Competitive risk from Booking.com and Airbnb "
        "is real; we do not claim we will defeat them globally. We claim a focused niche: "
        "fairer economics for independents plus planner-to-book integration.",
    )

    # 14
    h(doc, "14. Future Work After Capstone", 1)
    p(
        doc,
        "If the team continues HiddenStay beyond the course, the next practical steps are: "
        "upgrade hosting so the API stays awake; replace demo login with real authentication; "
        "onboard a small set of real pilot hosts in one city; measure whether the planner "
        "increases booking intent; and only then expand features such as richer payouts, more "
        "cities, or paid Growth tools after proof. The product should stay lean. The goal is "
        "survival and learning, not pretending we already operate at OTA scale.",
    )

    # 15
    h(doc, "15. Conclusion", 1)
    p(
        doc,
        "HiddenStay exists because high OTA commissions hurt independent hosts, and because "
        "travellers deserve a smoother path from trip idea to booked stay. Our solution is a "
        "fair-commission marketplace with a free host portal and an AI planner that sits on "
        "top of real inventory. We have implemented the core MVP and deployed a live demo for "
        "capstone evaluation. The money model is transparent: from 5% on bookings, optional "
        "upgrades, and a fairness cap. The technology stack is modern and suitable for a "
        "student-built MVP, with known hosting limits we can upgrade later.",
    )
    p(
        doc,
        "Final one-sentence summary for any reader: HiddenStay helps Southeast Asia homestay "
        "hosts keep 95% of a booking, and helps travellers plan and book authentic stays in "
        "one place.",
        bold=True,
    )

    # 16
    h(doc, "16. Related Documents", 1)
    p(
        doc,
        "Readers who want more detail can open these team documents in the project repository:",
    )
    bullets(
        doc,
        [
            "marketing/PROJECT_BRIEF.md — full product brief",
            "marketing/HIDDENSTAY_MONEY_GUIDE.md — money model in plain English",
            "marketing/CEO_LOUNGE_ONE_PAGER.md — short executive one-pager",
            "marketing/HiddenStay_Poster_Content_Brief.docx — poster designer brief",
            "README.md — technical quick start and live links",
            "DEPLOYMENT.md — how the live system is hosted",
        ],
    )

    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    er = end.add_run(
        "— End of Project Proposal —\n"
        "HiddenStay · Group 63 I · James Cook University Singapore · July 2026\n"
        "Any financial figures in related money documents are planning estimates, not audited statements."
    )
    style_run(er, size=10, italic=True)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
