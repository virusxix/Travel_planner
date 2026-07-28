"""Generate a simple, clear HiddenStay poster brief DOCX."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(r"D:\Development\Projects\sample\marketing\HiddenStay_Poster_Content_Brief.docx")


def h(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x15, 0x10)


def p(doc, text, bold=False, italic=False, size=12):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    return para


def bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(item, style="List Bullet")
        for run in para.runs:
            run.font.size = Pt(12)
            run.font.name = "Calibri"


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = header
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.size = Pt(11)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.rows[r_i + 1].cells[c_i]
            cell.text = val
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(11)
    doc.add_paragraph()


def main():
    doc = Document()
    s = doc.sections[0]
    s.top_margin = Inches(0.85)
    s.bottom_margin = Inches(0.85)
    s.left_margin = Inches(0.95)
    s.right_margin = Inches(0.95)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("HiddenStay Poster Brief")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0xC4, 0x5C, 0x3E)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Simple guide for the poster designer · Group 63 I · Jul 2026")
    sr.font.size = Pt(12)
    sr.italic = True

    p(
        doc,
        "Use this file as the content checklist. Put the words below on the poster. "
        "You choose the design look.",
        size=12,
    )

    # What is it
    h(doc, "1. What is HiddenStay? (one line)", 1)
    p(
        doc,
        "A booking app for small homestays in Southeast Asia. Hosts keep 95%. "
        "Guests can plan a trip with AI and book a stay in the same app.",
        bold=True,
        size=13,
    )

    # Print rules
    h(doc, "2. Print rules (must follow)", 1)
    bullets(
        doc,
        [
            "Size: A1 portrait (tall poster)",
            "Leave empty space near the edges (about 1.5 inches / 38 mm)",
            "Big title: HiddenStay",
            "Left logo: JCUS seal",
            "Right logo: James Cook University Singapore",
            "Save as PDF, high quality (200–300 DPI)",
        ],
    )

    # Layout order
    h(doc, "3. Layout order (top → bottom)", 1)
    p(doc, "People read from top to bottom. Put sections in this order:", size=12)
    bullets(
        doc,
        [
            "Logos + HiddenStay title",
            "The problem (big numbers)",
            "Hook line",
            "What we built (product + phone pictures)",
            "How we make money",
            "Why us (short)",
            "QR code + team names at the bottom",
        ],
    )

    # Exact text
    h(doc, "4. Exact text to put on the poster", 1)

    h(doc, "Title area", 2)
    bullets(
        doc,
        [
            "Title: HiddenStay",
            "Under title: Fair Commission Homestays + AI Trip Planner · Southeast Asia",
        ],
    )

    h(doc, "The problem", 2)
    p(
        doc,
        "Independent homestays lose 15–30% per booking to Booking.com / Expedia. "
        "Travellers plan trips in one place and book in another.",
    )
    table(
        doc,
        ["Show this big number", "Meaning"],
        [
            ["15–30%", "What OTAs take today"],
            ["5%", "What HiddenStay takes"],
            ["95%", "What the host keeps"],
        ],
    )
    p(doc, "Small note: On a SGD 100 stay, host saves SGD 13+ vs an 18% OTA fee.", italic=True)

    h(doc, "Hook line (make this big)", 2)
    p(doc, "What if hosts could", size=12)
    p(doc, "List free · Keep 95% · Plan & book in one app?", bold=True, size=14)

    h(doc, "What we built (MVP)", 2)
    bullets(
        doc,
        [
            "10-week capstone · Group 63 I",
            "5 pilot listings",
            "AI trip planner + map",
            "Booking with Stripe (test mode)",
            "Free host portal (list, bookings, earnings)",
        ],
    )

    h(doc, "Mission & vision (short)", 2)
    p(doc, "Mission: Fair fees for hosts + easy plan → map → book for travellers.", bold=False)
    p(doc, "Vision: Trusted booking for hidden stays in SEA — one city at a time.")

    h(doc, "Three phone screens (pictures)", 2)
    p(doc, "If you can, use real screenshots from the live app. Labels:")
    table(
        doc,
        ["Phone", "Show"],
        [
            ["1. AI Planner", "Trip plan chat (example: Chiang Mai)"],
            ["2. Map + Book", "Map with stays + Book button"],
            ["3. Host Portal", "Earnings screen (host keeps 95%)"],
        ],
    )

    h(doc, "How we make money", 2)
    p(doc, "Show these 4 lines clearly:")
    table(
        doc,
        ["Stream", "What hosts pay", "Required?"],
        [
            ["1. Booking fee", "5% when a guest books", "Yes"],
            ["2. Growth tools", "SGD 50 / month", "No — optional"],
            ["3. Featured listing", "SGD 99 / month", "No — optional"],
            ["4. Ads", "7–10% only on bookings we bring", "No — optional · never over 15% total"],
        ],
    )
    p(
        doc,
        "Important: Host portal is free. Do not write “only 5%” without saying upgrades are optional.",
        bold=True,
        size=11,
    )

    h(doc, "Why us (keep very short)", 2)
    table(
        doc,
        ["", "Write only this"],
        [
            ["Good", "Lower fee than OTAs · Free portal · Plan and book together"],
            ["Hard", "New brand · 5% alone is small money per booking"],
            ["Chance", "Many small homestays in SEA"],
            ["Risk", "Big OTAs already well known"],
        ],
    )

    h(doc, "Feature chips (7 short labels)", 2)
    bullets(
        doc,
        [
            "5% marketplace",
            "AI planner + map",
            "Plan → book",
            "Free host portal",
            "Growth tools",
            "15% fee cap",
            "18-month plan",
        ],
    )

    h(doc, "QR + call to action", 2)
    bullets(
        doc,
        [
            "QR code opens: https://hidden-stay.vercel.app",
            "Text under QR: Scan to try the live demo",
            "QR image file: marketing/hiddenstay-qr.png",
        ],
    )

    h(doc, "Bottom / team", 2)
    p(doc, "BU3102 / CP3102 Capstone · Group 63 I · Jul 2026", bold=True)
    table(
        doc,
        ["Role", "Name"],
        [
            ["PM", "Zaw Latt Naung"],
            ["Dev", "Oakkar Phyoe"],
            ["Design / UAT", "Eka Dian Tara Tiu"],
            ["BA / AI", "Chenyang Gu"],
            ["Marketing", "Yihua Deng"],
        ],
    )
    p(
        doc,
        "Small footer note: Money figures are estimates for planning — not audited accounts.",
        italic=True,
        size=10,
    )

    # Do / don't
    h(doc, "5. Do / Don’t", 1)
    h(doc, "Do", 2)
    bullets(
        doc,
        [
            "Use big numbers people can read from far away",
            "Keep text short",
            "Show the QR clearly",
            "Use warm orange / terracotta accent (like the app)",
        ],
    )
    h(doc, "Don’t", 2)
    bullets(
        doc,
        [
            "Don’t invent user counts or revenue",
            "Don’t write long paragraphs",
            "Don’t say we beat Booking.com worldwide",
            "Don’t put code / tech stack on the poster",
            "Don’t hide the optional fees if you show “5%”",
        ],
    )

    # Deliver
    h(doc, "6. What to send back", 1)
    bullets(
        doc,
        [
            "A1 PDF ready to print",
            "Editable file (Figma / Canva / PowerPoint / Illustrator — any is fine)",
            "Check: logos are real JCU logos",
            "Check: QR opens the live website",
        ],
    )

    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    er = end.add_run("Questions? Ask the team PM. — HiddenStay Group 63 I")
    er.italic = True
    er.font.size = Pt(11)

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
