"""Fill official JCU A1 poster template with HiddenStay content. Logos stay fixed."""
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

SRC = Path(r"d:\c\Downloads\pwh 38TH-Poster-template corrent.pptx")
OUT = Path(r"D:\Development\Projects\sample\marketing\HiddenStay-JCU-Poster.pptx")
OUT2 = Path(r"d:\c\Downloads\HiddenStay-JCU-Poster.pptx")

CREAM = RGBColor(0xF5, 0xED, 0xE0)
PANEL = RGBColor(0xE8, 0xDC, 0xC8)
TAN = RGBColor(0xD4, 0xC4, 0xA8)
BROWN = RGBColor(0x5C, 0x40, 0x33)
DARK = RGBColor(0x3D, 0x2E, 0x24)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
STAT = RGBColor(0x8B, 0x3A, 0x2A)
GREEN = RGBColor(0x2D, 0x6A, 0x5A)
MUTED = RGBColor(0x6B, 0x5A, 0x48)
INK = RGBColor(0x1A, 0x15, 0x10)
SWOT_S = RGBColor(0x5B, 0x7C, 0x99)
SWOT_W = RGBColor(0xD4, 0xA8, 0x43)
SWOT_O = RGBColor(0xD4, 0x84, 0x5B)
SWOT_T = RGBColor(0x8B, 0x6B, 0x9E)
ORANGE = RGBColor(0xC4, 0x7A, 0x3A)


def solid(shape, fill, line=None, lw=1.25):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(lw)


def box(slide, l, t, w, h, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(2)
        r = p.runs[0]
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return tb


def panel(slide, l, t, w, h, fill=PANEL):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    solid(s, fill, BROWN, 1.5)
    try:
        s.adjustments[0] = 0.05
    except Exception:
        pass
    return s


def hdr(slide, l, t, w, label):
    box(slide, l, t, w, Inches(0.35), label, size=18, bold=True, color=BROWN, font="Arial Black")


def clear_placeholder_text(slide):
    """Remove instruction body; keep logos + margin rect; rewrite title/refs."""
    kill = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            joined = "\n".join(p.text for p in sh.text_frame.paragraphs)
            if "A1 paper dimensions" in joined or "ALL POSTERS MUST BE VERTICAL" in joined:
                kill.append(sh)
            elif "Please your TITLE" in joined or "60 points" == joined.strip():
                # rewrite title in place
                tf = sh.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.CENTER
                r = p.add_run()
                r.text = "HiddenStay"
                r.font.size = Pt(60)
                r.font.bold = True
                r.font.color.rgb = BROWN
                r.font.name = "Times New Roman"
                p2 = tf.add_paragraph()
                p2.alignment = PP_ALIGN.CENTER
                r2 = p2.add_run()
                r2.text = "Fair commission homestays + AI trip planner · Southeast Asia"
                r2.font.size = Pt(16)
                r2.font.bold = True
                r2.font.color.rgb = MUTED
                r2.font.name = "Calibri"
            elif "Place your reference" in joined:
                tf = sh.text_frame
                tf.clear()
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT
                r = p.add_run()
                r.text = (
                    "Group 63 I — Zaw Latt Naung · Oakkar Phyoe · Eka Dian Tara Tiu · Chenyang Gu · Yihua Deng  |  "
                    "BU3102 / CP3102 Capstone · JCUS  |  Break-even Month 18 · 55 hosts · 92 bookings/mo  |  hiddenstay.app"
                )
                r.font.size = Pt(10)
                r.font.color.rgb = MUTED
                r.font.name = "Calibri"
        if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and sh.name.startswith("Rectangle"):
            solid(sh, CREAM, None)
    for sh in kill:
        el = sh._element
        el.getparent().remove(el)


def main():
    prs = Presentation(str(SRC))
    slide = prs.slides[0]
    clear_placeholder_text(slide)

    # Content band inside ~1.5" margins, below logos (~5.6") and above refs (~31.5")
    m = Inches(1.5)
    x = m
    y = Inches(5.55)
    right = prs.slide_width - m
    bottom = Inches(31.35)
    cw = right - x
    gap = Inches(0.18)
    col = (cw - gap) / 2

    # ── TOP ROW: Issue | Hook | MVP ──
    top_h = Inches(3.35)
    issue_w = col * 0.95
    hook_w = Inches(4.2)
    mvp_w = cw - issue_w - hook_w - 2 * gap

    panel(slide, x, y, issue_w, top_h)
    hdr(slide, x + Inches(0.2), y + Inches(0.15), issue_w - Inches(0.4), "THE ISSUE")
    box(
        slide,
        x + Inches(0.2),
        y + Inches(0.55),
        issue_w - Inches(0.4),
        Inches(0.7),
        "Independent homestays lose 15–30% per booking to OTAs. Travellers plan on Google and book on different apps.",
        size=15,
        color=DARK,
    )
    stats = [("15–30%", "OTA fee today"), ("5%", "HiddenStay fee"), ("95%", "Host keeps")]
    sw = (issue_w - Inches(0.6)) / 3
    for i, (num, lbl) in enumerate(stats):
        sx = x + Inches(0.2) + i * (sw + Inches(0.1))
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, sx, y + Inches(1.4), sw, Inches(1.45))
        solid(card, TAN, BROWN, 1)
        box(slide, sx, y + Inches(1.5), sw, Inches(0.85), num, size=32, bold=True, color=STAT, align=PP_ALIGN.CENTER, font="Arial Black", anchor=MSO_ANCHOR.MIDDLE)
        box(slide, sx, y + Inches(2.35), sw, Inches(0.4), lbl.upper(), size=11, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
    box(
        slide,
        x + Inches(0.2),
        y + Inches(2.95),
        issue_w - Inches(0.4),
        Inches(0.3),
        "SGD 100 stay → host saves SGD 13+ vs 18% OTA",
        size=13,
        bold=True,
        color=DARK,
    )

    hx = x + issue_w + gap
    hook = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, hx, y, hook_w, top_h)
    solid(hook, BROWN, BROWN, 1.5)
    box(slide, hx + Inches(0.2), y + Inches(0.45), hook_w - Inches(0.4), Inches(0.4), "WHAT IF HOSTS COULD", size=14, bold=True, color=RGBColor(0xE8, 0xDC, 0xC8), align=PP_ALIGN.CENTER)
    box(
        slide,
        hx + Inches(0.25),
        y + Inches(1.0),
        hook_w - Inches(0.5),
        Inches(1.5),
        "List free\nKeep 95%\nPlan & book\nin one app?",
        size=26,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
        font="Times New Roman",
    )
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, hx + hook_w / 2 - Inches(0.4), y + Inches(2.6), Inches(0.8), Inches(0.8))
    solid(circ, GREEN, WHITE, 2)
    box(slide, hx + hook_w / 2 - Inches(0.4), y + Inches(2.75), Inches(0.8), Inches(0.5), "HS", size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    mx = hx + hook_w + gap
    panel(slide, mx, y, mvp_w, top_h)
    hdr(slide, mx + Inches(0.2), y + Inches(0.15), mvp_w - Inches(0.4), "CAPSTONE MVP")
    for i, line in enumerate([
        "10-week build · Group 63 I",
        "5 pilot listings",
        "Stripe sandbox · AI planner",
        "≥80% map-valid stops",
        "Host portal free (Starter)",
    ]):
        box(slide, mx + Inches(0.25), y + Inches(0.6) + i * Inches(0.45), mvp_w - Inches(0.5), Inches(0.4), f"•  {line}", size=15, color=DARK)

    y2 = y + top_h + gap
    mid_h = Inches(4.0)

    # ── MID: Mission/phones | Business + SWOT ──
    panel(slide, x, y2, col, mid_h)
    # mission / vision
    mw = (col - Inches(0.55)) / 2
    for i, (title, body) in enumerate([
        ("MISSION", "Fair host economics + seamless plan → map → book for travellers."),
        ("VISION", "Trusted direct-booking for hidden stays across SEA — one city at a time."),
    ]):
        bx = x + Inches(0.18) + i * (mw + Inches(0.18))
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, y2 + Inches(0.18), mw, Inches(1.35))
        solid(card, TAN, BROWN, 1)
        box(slide, bx + Inches(0.12), y2 + Inches(0.28), mw - Inches(0.24), Inches(0.3), title, size=12, bold=True, color=BROWN, font="Arial Black")
        box(slide, bx + Inches(0.12), y2 + Inches(0.6), mw - Inches(0.24), Inches(0.8), body, size=13, color=DARK)

    # phone mockups
    phones = [
        ("AI PLANNER", "Day 1 · Chiang Mai\n● Old town\n● Homestay\nDay 2 · countryside"),
        ("MAP + BOOK", "[ MAP ]\n● Stay A  SGD 50/n\n● Stay B  SGD 45/n\n→ Book"),
        ("HOST PORTAL", "Earnings\nAvail. SGD 285\nPending SGD 95\n5% · 95% yours"),
    ]
    pw, ph = Inches(2.35), Inches(4.15)  # will scale down
    pw, ph = Inches(2.2), Inches(2.15)
    total = 3 * pw + 2 * Inches(0.15)
    px0 = x + (col - total) / 2
    py0 = y2 + Inches(1.7)
    for i, (lbl, body) in enumerate(phones):
        px = px0 + i * (pw + Inches(0.15))
        phone = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, py0, pw, ph)
        solid(phone, INK, BROWN, 1.5)
        box(slide, px, py0 + Inches(0.08), pw, Inches(0.28), lbl, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        screen = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px + Inches(0.1), py0 + Inches(0.4), pw - Inches(0.2), ph - Inches(0.5))
        solid(screen, RGBColor(0x0F, 0x2A, 0x3A), None)
        box(slide, px + Inches(0.15), py0 + Inches(0.5), pw - Inches(0.3), ph - Inches(0.65), body, size=10, color=RGBColor(0xA8, 0xE6, 0xCF), align=PP_ALIGN.CENTER)

    rx = x + col + gap
    panel(slide, rx, y2, col, mid_h)
    hdr(slide, rx + Inches(0.2), y2 + Inches(0.12), col - Inches(0.4), "BUSINESS MODEL")
    nodes = [("FOCUS", "Supply + planner"), ("STARTER", "Free · 5%"), ("REVENUE", "4 streams")]
    nw = (col - Inches(0.8)) / 3
    for i, (a, b) in enumerate(nodes):
        nx = rx + Inches(0.2) + i * (nw + Inches(0.1))
        n = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, nx, y2 + Inches(0.55), nw, Inches(0.85))
        solid(n, TAN, BROWN, 1)
        box(slide, nx, y2 + Inches(0.6), nw, Inches(0.35), a, size=12, bold=True, color=BROWN, align=PP_ALIGN.CENTER, font="Arial Black")
        box(slide, nx, y2 + Inches(0.95), nw, Inches(0.35), b, size=12, color=DARK, align=PP_ALIGN.CENTER)
    box(
        slide,
        rx + Inches(0.2),
        y2 + Inches(1.5),
        col - Inches(0.4),
        Inches(0.45),
        "① 5% booking   ② Growth SGD 50/mo   ③ Featured SGD 99   ④ Ads 7–10% (15% cap)",
        size=12,
        bold=True,
        color=DARK,
    )

    hdr(slide, rx + Inches(0.2), y2 + Inches(2.0), col - Inches(0.4), "WHY US? (SWOT)")
    swot = [
        (SWOT_S, "S", "5% vs 15–30% OTAs · Free portal · Plan→book"),
        (SWOT_W, "W", "Thin unit economics · New brand"),
        (SWOT_O, "O", "SEA homestay supply · SEO + planner"),
        (SWOT_T, "T", "OTA dominance · Guest discovery cost"),
    ]
    for i, (colr, letter, body) in enumerate(swot):
        sy = y2 + Inches(2.4) + i * Inches(0.38)
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, rx + Inches(0.2), sy, Inches(0.45), Inches(0.32))
        solid(badge, colr, colr)
        box(slide, rx + Inches(0.2), sy, Inches(0.45), Inches(0.32), letter, size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        box(slide, rx + Inches(0.75), sy, col - Inches(1.0), Inches(0.32), body, size=12, color=DARK, anchor=MSO_ANCHOR.MIDDLE)

    # ── BOTTOM: features ──
    y3 = y2 + mid_h + gap
    bot_h = bottom - y3
    panel(slide, x, y3, cw, bot_h)
    hdr(slide, x + Inches(0.25), y3 + Inches(0.15), cw - Inches(0.5), "PRODUCT + BUSINESS FEATURES")
    feats = [
        "5% marketplace",
        "AI planner + map",
        "Plan → book flow",
        "Free host portal",
        "Growth tools (SEO)",
        "Multi-night trips",
        "18-mo survival plan",
    ]
    fw = (cw - Inches(0.7)) / 7
    for i, f in enumerate(feats):
        fx = x + Inches(0.25) + i * (fw + Inches(0.05))
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, fx, y3 + Inches(0.6), fw, bot_h - Inches(0.85))
        solid(card, TAN, BROWN, 1)
        num = slide.shapes.add_shape(MSO_SHAPE.OVAL, fx + fw / 2 - Inches(0.28), y3 + Inches(0.85), Inches(0.56), Inches(0.56))
        solid(num, BROWN, None)
        box(slide, fx + fw / 2 - Inches(0.28), y3 + Inches(0.95), Inches(0.56), Inches(0.4), str(i + 1), size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        box(slide, fx + Inches(0.08), y3 + Inches(1.55), fw - Inches(0.16), Inches(1.2), f.upper(), size=13, bold=True, color=DARK, align=PP_ALIGN.CENTER)

    prs.save(str(OUT))
    prs.save(str(OUT2))
    print("Saved", OUT)
    print("Saved", OUT2)
    # quick checks
    assert OUT.stat().st_size > 100_000
    pics = sum(1 for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert pics == 2, pics
    print("OK logos=", pics)


if __name__ == "__main__":
    main()
